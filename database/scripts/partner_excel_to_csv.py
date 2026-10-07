"""제휴매장 엑셀의 '식당별 요약' 시트를 DB 적재용 CSV로 바꾼다.

사용법:
  python partner_excel_to_csv.py 제휴매장.xlsx -o 제휴혜택.csv

출력 한 행은 한 가게의 한 제휴 혜택이다. 여러 단과대가 같은 혜택을 받으면
colleges 열에 | 로 연결해 저장한다.
"""

import argparse
import csv
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("openpyxl 이 필요합니다: pip install openpyxl")


COLLEGE_NAMES = {
    "정책법학대학",
    "전자정보공과대학",
    "자연과학대학",
    "인문사회과학대학",
    "공과대학",
    "경영대학",
    "인공지능융합대학",
}


def text(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def benefits(value):
    raw = str(value or "").strip()
    pattern = re.compile(r"\[([^\]]+)]\s*([\s\S]*?)(?=\n\s*\[|$)")
    return [(colleges.strip(), text(offer)) for colleges, offer in pattern.findall(raw)]


def convert(input_path, output_path):
    workbook = load_workbook(input_path, data_only=True, read_only=True)
    if "식당별 요약" not in workbook.sheetnames:
        sys.exit("'식당별 요약' 시트를 찾을 수 없습니다.")
    sheet = workbook["식당별 요약"]
    headers = {text(cell.value): index for index, cell in enumerate(next(sheet.iter_rows()), start=1)}
    required = {"식당명", "단과대학", "제휴 내용"}
    missing = required - headers.keys()
    if missing:
        sys.exit(f"필수 열이 없습니다: {', '.join(sorted(missing))}")

    rows = []
    errors = []
    for source_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        store = text(values[headers["식당명"] - 1])
        college_text = text(values[headers["단과대학"] - 1])
        offer_text = values[headers["제휴 내용"] - 1]
        if not store and not college_text and not text(offer_text):
            continue
        if store == "합계":
            continue
        expected = {name.strip() for name in college_text.split(",") if name.strip()}
        parsed = benefits(offer_text)
        if not parsed:
            errors.append(f"{source_row}행 {store}: [단과대학] 형식의 제휴 내용을 찾지 못함")
            continue
        covered = set()
        for sequence, (college_group, offer) in enumerate(parsed, start=1):
            colleges = [name.strip() for name in college_group.split("·") if name.strip()]
            unknown = set(colleges) - COLLEGE_NAMES
            if unknown:
                errors.append(f"{source_row}행 {store}: 알 수 없는 단과대 {', '.join(sorted(unknown))}")
            if not offer:
                errors.append(f"{source_row}행 {store}: 빈 혜택")
            covered.update(colleges)
            rows.append({
                "store_name": store,
                "colleges": "|".join(colleges),
                "offer": offer,
                "source_ref": f"summary-{source_row}-{sequence}",
                "source_row": source_row,
                "data_source": input_path.name,
            })
        if expected != covered:
            missing_colleges = expected - covered
            extra_colleges = covered - expected
            detail = []
            if missing_colleges:
                detail.append(f"혜택 없는 단과대 {', '.join(sorted(missing_colleges))}")
            if extra_colleges:
                detail.append(f"요약에 없는 단과대 {', '.join(sorted(extra_colleges))}")
            errors.append(f"{source_row}행 {store}: {'; '.join(detail)}")

    if errors:
        sys.exit("제휴 자료 검증 실패:\n  " + "\n  ".join(errors))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8-sig") as output:
        writer = csv.DictWriter(output, fieldnames=[
            "store_name", "colleges", "offer", "source_ref", "source_row", "data_source"
        ])
        writer.writeheader()
        writer.writerows(rows)

    print(f"저장: {output_path}")
    print(f"혜택 {len(rows)}개 / 가게 {len({row['store_name'] for row in rows})}곳")


def main():
    parser = argparse.ArgumentParser(description="제휴매장 엑셀 -> 제휴 혜택 CSV")
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("-o", "--output", type=Path, required=True)
    args = parser.parse_args()
    if not args.xlsx.exists():
        sys.exit(f"파일을 찾을 수 없습니다: {args.xlsx}")
    convert(args.xlsx, args.output)


if __name__ == "__main__":
    main()
